from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable

from app.models import SourceDocument

SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".pdf", ".docx"}


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".csv":
        with path.open(encoding="utf-8", errors="ignore", newline="") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle))
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    raise ValueError(f"不支持的文档格式：{suffix}")


def load_documents(paths: Iterable[Path]) -> list[SourceDocument]:
    documents: list[SourceDocument] = []
    for path in paths:
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        content = read_document(path)
        if content.strip():
            documents.append(SourceDocument(
                title=path.stem,
                content=content,
                snippet=re.sub(r"\s+", " ", content)[:240],
                source_type="local",
                domain="local-document",
                metadata={"path": str(path)},
            ))
    return documents


def load_and_chunk_documents(
    paths: Iterable[Path], chunk_size: int = 900, overlap: int = 120
) -> list[SourceDocument]:
    chunks: list[SourceDocument] = []
    for document in load_documents(paths):
        chunks.extend(chunk_document(document, chunk_size=chunk_size, overlap=overlap))
    return chunks


def chunk_document(document: SourceDocument, chunk_size: int = 900, overlap: int = 120) -> list[SourceDocument]:
    """Split by headings/paragraphs first, then enforce bounded chunks."""
    blocks = [block.strip() for block in re.split(r"\n(?=#{1,6}\s)|\n\s*\n", document.content) if block.strip()]
    chunks: list[str] = []
    current = ""
    for block in blocks:
        if len(current) + len(block) + 2 <= chunk_size:
            current = f"{current}\n\n{block}".strip()
            continue
        if current:
            chunks.append(current)
        if len(block) <= chunk_size:
            current = block
        else:
            start = 0
            while start < len(block):
                chunks.append(block[start:start + chunk_size])
                start += max(1, chunk_size - overlap)
            current = ""
    if current:
        chunks.append(current)

    result: list[SourceDocument] = []
    for index, chunk in enumerate(chunks):
        item = document.model_copy(deep=True)
        item.title = f"{document.title} · 片段 {index + 1}"
        item.content = chunk
        item.snippet = re.sub(r"\s+", " ", chunk)[:240]
        item.metadata = {
            **document.metadata,
            "parent_title": document.title,
            "chunk_index": index,
            "chunk_count": len(chunks),
        }
        result.append(item)
    return result
