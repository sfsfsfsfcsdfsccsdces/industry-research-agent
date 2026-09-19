from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _plain_text(markdown: str) -> str:
    text = re.sub(r"```.*?```", "", markdown, flags=re.DOTALL)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"[*_`>]", "", text)
    return text


class ReportExporter:
    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_all(self, task_id: str, markdown: str) -> dict[str, str]:
        target = self.output_dir / task_id
        target.mkdir(parents=True, exist_ok=True)
        md_path = target / "report.md"
        docx_path = target / "report.docx"
        pdf_path = target / "report.pdf"
        md_path.write_text(markdown, encoding="utf-8")
        self._to_docx(markdown, docx_path)
        self._to_pdf(markdown, pdf_path)
        return {"markdown": str(md_path), "docx": str(docx_path), "pdf": str(pdf_path)}

    def _to_docx(self, markdown: str, path: Path) -> None:
        document = Document()
        for line in markdown.splitlines():
            clean = _plain_text(line).strip()
            if not clean:
                continue
            if line.startswith("# "):
                document.add_heading(clean.lstrip("# "), level=0)
            elif line.startswith("## "):
                document.add_heading(clean.lstrip("# "), level=1)
            elif line.startswith("### "):
                document.add_heading(clean.lstrip("# "), level=2)
            elif line.startswith("- "):
                document.add_paragraph(clean[2:], style="List Bullet")
            else:
                document.add_paragraph(clean)
        document.save(path)

    def _to_pdf(self, markdown: str, path: Path) -> None:
        font_name = "STSong-Light"
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
        for candidate in (
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        ):
            if Path(candidate).exists():
                try:
                    pdfmetrics.registerFont(TTFont("ReportFont", candidate))
                    font_name = "ReportFont"
                    break
                except Exception:
                    continue
        styles = getSampleStyleSheet()
        body = ParagraphStyle("ChineseBody", parent=styles["BodyText"], fontName=font_name, fontSize=9.5, leading=15, spaceAfter=6)
        title = ParagraphStyle("ChineseTitle", parent=body, fontSize=18, leading=25, alignment=TA_CENTER, spaceAfter=14)
        heading = ParagraphStyle("ChineseHeading", parent=body, fontSize=13, leading=19, spaceBefore=8, spaceAfter=6)
        story = []
        for line in markdown.splitlines():
            clean = _plain_text(line).strip()
            if not clean or clean.startswith("|---"):
                continue
            clean = clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            style = title if line.startswith("# ") else heading if line.startswith("##") else body
            story.append(Paragraph(clean.lstrip("# "), style))
            story.append(Spacer(1, 1.5 * mm))
        doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm)
        doc.build(story)
